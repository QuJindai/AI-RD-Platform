"""Bounded, nonexecuting source extraction for the knowledge library.

PDF and XLSX parsers run in a short-lived process, with process resource limits
where supported. No macros, HTML scripts, formulas or document links execute.
"""
from datetime import date, datetime
from html.parser import HTMLParser
from io import BytesIO
import json
import os
from pathlib import PurePosixPath
import re
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET


MAX_SOURCE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
MAX_ZIP_BYTES = 32 * 1024 * 1024
MAX_ZIP_MEMBERS = 1000
MAX_ZIP_RATIO = 200
MAX_PDF_PAGES = 300
MAX_PDF_STREAM_BYTES = 8 * 1024 * 1024
MAX_PDF_TOTAL_STREAM_BYTES = 32 * 1024 * 1024
MAX_SHEET_CELLS = 100_000
MAX_SEGMENTS = 100_000
WORKER_TIMEOUT = 15


class TextBuilder:
    def __init__(self):
        self.parts, self.segments, self.length = [], [], 0

    def append(self, text, location, separator='\n'):
        if not text:
            return
        prefix = separator if self.parts else ''
        if self.length + len(prefix) + len(text) > MAX_TEXT_CHARS:
            raise ValueError('提取文本超过200万字符限制')
        if len(self.segments) >= MAX_SEGMENTS:
            raise ValueError('文档来源位置数量超过限制')
        self.parts.extend((prefix, text))
        start = self.length + len(prefix)
        self.length = start + len(text)
        self.segments.append({'start': start, 'end': self.length, **location})

    def result(self):
        text = ''.join(self.parts)
        if not text.strip():
            raise ValueError('未提取到可检索文本；扫描PDF需要先进行OCR')
        return {'text': text, 'segments': self.segments}


def _decode(raw):
    try:
        text = raw.decode('utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')
    except UnicodeError as exc:
        raise ValueError('文本文件需要UTF-8或带BOM的UTF-16编码') from exc
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError('提取文本超过200万字符限制')
    if '\x00' in text:
        raise ValueError('文本文件包含空字节，可能为二进制内容')
    return text


def _safe_zip(raw):
    try:
        archive = zipfile.ZipFile(BytesIO(raw))
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            raise ValueError('压缩文档文件数量超过限制')
        total, names = 0, set()
        for member in members:
            path = PurePosixPath(member.filename)
            if (path.is_absolute() or '..' in path.parts or '\\' in member.filename
                    or ':' in member.filename or member.filename in names
                    or (member.external_attr >> 16) & 0o170000 == 0o120000):
                raise ValueError('压缩文档包含不安全或重复路径')
            names.add(member.filename)
            if member.flag_bits & 1:
                raise ValueError('不支持加密压缩文档')
            total += member.file_size
            if (total > MAX_ZIP_BYTES or member.file_size > MAX_ZIP_BYTES
                    or member.file_size > max(member.compress_size, 1) * MAX_ZIP_RATIO):
                raise ValueError('压缩文档解压大小或压缩比超过限制')
        # Read through the bounded ZIP reader before handing XML to a parser.
        for member in members:
            if member.filename.lower().endswith(('.xml', '.rels')):
                content = archive.read(member)
                # DTD and entity declarations are unnecessary in OOXML.
                probe = content.replace(b'\x00', b'').upper()
                if b'<!DOCTYPE' in probe or b'<!ENTITY' in probe:
                    raise ValueError('压缩文档禁止DTD或实体声明')
        return archive
    except (zipfile.BadZipFile, RuntimeError, NotImplementedError, OSError) as exc:
        raise ValueError('压缩文档损坏或格式不支持') from exc


def _docx(raw):
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with _safe_zip(raw) as archive:
        if 'word/document.xml' not in archive.namelist():
            raise ValueError('DOCX缺少word/document.xml')
        names = ['word/document.xml'] + sorted(name for name in archive.namelist()
                    if re.fullmatch(r'word/(?:header\d+|footer\d+|footnotes|endnotes)\.xml', name))
        builder = TextBuilder()
        try:
            for part in names:
                tree = ET.fromstring(archive.read(part))
                for index, paragraph in enumerate(tree.iter(namespace + 'p'), 1):
                    pieces = []
                    for element in paragraph.iter():
                        if element.tag == namespace + 't':
                            pieces.append(element.text or '')
                        elif element.tag in (namespace + 'br', namespace + 'cr'):
                            pieces.append('\n')
                        elif element.tag == namespace + 'tab':
                            pieces.append('\t')
                    builder.append(''.join(pieces), {'part': part, 'paragraph': index}, '\n\n')
        except ET.ParseError as exc:
            raise ValueError('DOCX XML损坏') from exc
        return builder.result()


class _HTMLText(HTMLParser):
    ignored = {'script', 'style', 'template', 'noscript', 'head', 'iframe', 'object'}
    blocks = {'p', 'div', 'section', 'article', 'li', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6',
              'br', 'tr', 'table', 'ul', 'ol', 'blockquote', 'pre', 'hr'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.builder, self.hidden, self.buffer = TextBuilder(), [], []
        self.block, self.buffer_size = 0, 0

    def flush(self):
        value = ''.join(self.buffer).strip()
        self.buffer, self.buffer_size = [], 0
        if value:
            self.block += 1
            self.builder.append(value, {'html_block': self.block}, '\n\n')

    def handle_starttag(self, tag, attrs):
        if tag in self.ignored:
            self.hidden.append(tag)
        if not self.hidden and tag in self.blocks:
            self.flush()

    def handle_endtag(self, tag):
        if self.hidden:
            if tag == self.hidden[-1]:
                self.hidden.pop()
            return
        if tag in self.blocks:
            self.flush()

    def handle_data(self, data):
        if not self.hidden:
            self.buffer_size += len(data)
            if self.builder.length + self.buffer_size > MAX_TEXT_CHARS:
                raise ValueError('提取文本超过200万字符限制')
            self.buffer.append(data)


def _html(raw):
    parser = _HTMLText()
    parser.feed(_decode(raw))
    parser.close()
    parser.flush()
    return parser.builder.result()


def _pdf(raw):
    if not raw.startswith(b'%PDF-'):
        raise ValueError('PDF文件头无效')
    import pypdf
    from pypdf import filters
    # These globals are isolated to this disposable worker process.
    for name in ('ZLIB_MAX_OUTPUT_LENGTH', 'LZW_MAX_OUTPUT_LENGTH', 'RUN_LENGTH_MAX_OUTPUT_LENGTH',
                 'MAX_ARRAY_BASED_STREAM_OUTPUT_LENGTH', 'MAX_DECLARED_STREAM_LENGTH'):
        setattr(filters, name, MAX_PDF_STREAM_BYTES)
    reader = pypdf.PdfReader(BytesIO(raw), strict=True, root_object_recovery_limit=1000)
    if reader.is_encrypted:
        raise ValueError('不支持加密PDF')
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError('PDF超过300页限制')
    builder, content_size = TextBuilder(), 0
    for number, page in enumerate(reader.pages, 1):
        content = page.get_contents()
        if content is not None:
            content_size += len(content.get_data())
            if content_size > MAX_PDF_TOTAL_STREAM_BYTES:
                raise ValueError('PDF解压内容超过32MiB限制')
        value = page.extract_text() or ''
        builder.append(value, {'page': number}, '\n\n')
    return builder.result()


def _xlsx(raw):
    with _safe_zip(raw) as archive:
        if 'xl/workbook.xml' not in archive.namelist():
            raise ValueError('XLSX缺少工作簿')
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    builder, cells = TextBuilder(), 0
    try:
        if len(book.worksheets) > 100:
            raise ValueError('XLSX工作表超过100个限制')
        for sheet in book.worksheets:
            if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_SHEET_CELLS:
                raise ValueError('XLSX工作表声明范围超过10万单元格限制')
            for row in sheet.iter_rows():
                cells += len(row)
                if cells > MAX_SHEET_CELLS:
                    raise ValueError('XLSX超过10万单元格限制')
                values, coordinates = [], []
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else str(cell.value)
                    values.append(f'{cell.coordinate}: {value}')
                    coordinates.append(cell.coordinate)
                if values:
                    builder.append('\t'.join(values), {'sheet': sheet.title, 'cells': coordinates})
    finally:
        book.close()
    return builder.result()


def _worker(kind, raw):
    env = {**os.environ, 'OPENBLAS_NUM_THREADS': '1', 'OMP_NUM_THREADS': '1'}
    try:
        result = subprocess.run([sys.executable, '-m', 'ard.knowledge_parsers', kind], input=raw,
                                stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                                timeout=WORKER_TIMEOUT, env=env, check=False)
    except subprocess.TimeoutExpired as exc:
        raise ValueError('文档解析超过15秒限制') from exc
    if result.returncode or len(result.stdout) > 24 * 1024 * 1024:
        raise ValueError('文档解析失败或超过资源限制')
    try:
        response = json.loads(result.stdout)
    except (ValueError, UnicodeError) as exc:
        raise ValueError('文档解析失败') from exc
    if response.get('error'):
        raise ValueError(response['error'])
    return response


def parse_document(filename, raw):
    if not raw or len(raw) > MAX_SOURCE_BYTES:
        raise ValueError('导入文件必须为1字节至10MiB')
    if not isinstance(filename, str) or not filename or len(filename) > 255 or any(c in filename for c in ('/', '\\', '\x00', '\r', '\n')):
        raise ValueError('文件名无效')
    kind = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if kind in ('txt', 'md'):
        builder = TextBuilder()
        builder.append(_decode(raw), {'line_start': 1})
        result = builder.result()
    elif kind in ('html', 'htm'):
        result = _html(raw)
    elif kind in ('docx', 'pdf', 'xlsx'):
        result = _worker(kind, raw)
    else:
        raise ValueError('仅支持TXT、MD、DOCX、PDF、HTML、XLSX')
    return {**result, 'format': 'html' if kind == 'htm' else kind, 'filename': filename}


def _main():
    # POSIX limits supplement the cross-platform wall-clock timeout and byte bounds.
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
        resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
    except (ImportError, ValueError, OSError):
        pass
    try:
        raw = sys.stdin.buffer.read(MAX_SOURCE_BYTES + 1)
        if not raw or len(raw) > MAX_SOURCE_BYTES:
            raise ValueError('导入文件必须为1字节至10MiB')
        result = {'docx': _docx, 'pdf': _pdf, 'xlsx': _xlsx}[sys.argv[1]](raw)
    except ValueError as exc:
        result = {'error': str(exc)[:200]}
    except Exception:
        result = {'error': '文档损坏、格式不支持或超过解析资源限制'}
    sys.stdout.write(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    _main()
